import os
import base64
from io import BytesIO
from fastapi import APIRouter, UploadFile, File, Form, HTTPException
import docx
import openpyxl

from schemas_multimodal import MultimodalExtractResponse
from services.prompt_engine import prompt_engine
from services.llm import call_claude_vision_json

router = APIRouter(prefix="/multimodal", tags=["多模态解析"])


def _load_system_prompt(filename: str) -> str:
    """Load a raw prompt file for vision calls (can't use Jinja2 for vision API)."""
    path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "prompts", filename)
    with open(path, "r", encoding="utf-8") as f:
        return f.read()


@router.post(
    "/extract",
    response_model=MultimodalExtractResponse,
    summary="多模态抽取",
    description="从文本、图片或 Office 文件中抽取单词与释义。",
)
async def extract_vocabulary(
    file: UploadFile = File(...),
    domain: str = Form("general")
):
    contents = await file.read()
    filename = file.filename.lower()

    # ── Text-based file types: use PromptEngine ──
    extracted_text = ""

    if filename.endswith(".txt") or filename.endswith(".md"):
        extracted_text = contents.decode("utf-8")

    elif filename.endswith(".docx"):
        try:
            doc = docx.Document(BytesIO(contents))
            extracted_text = "\n".join([para.text for para in doc.paragraphs if para.text.strip() != ""])
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"解析 Word 文档失败：{str(e)}")

    elif filename.endswith(".xlsx"):
        try:
            wb = openpyxl.load_workbook(BytesIO(contents), data_only=True)
            sheet = wb.active
            rows = []
            for row in sheet.iter_rows(values_only=True):
                non_empty = [str(cell) for cell in row if cell is not None]
                if non_empty:
                    rows.append(" | ".join(non_empty))
            extracted_text = "\n".join(rows)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"解析 Excel 文档失败：{str(e)}")

    # ── Image files: use Vision API (can't use PromptEngine for vision) ──
    elif filename.endswith((".jpg", ".jpeg", ".png")):
        base64_image = base64.b64encode(contents).decode("utf-8")
        media_type = "image/jpeg" if filename.endswith(("jpg", "jpeg")) else "image/png"
        system_prompt = _load_system_prompt("multimodal.txt")
        user_prompt = f"Domain focus: {domain}\n\nPlease analyze the provided content."

        try:
            data = await call_claude_vision_json(
                system_prompt=system_prompt,
                user_prompt=user_prompt,
                base64_image=base64_image,
                media_type=media_type,
                max_tokens=4000
            )
            return MultimodalExtractResponse(**data)
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"视觉抽取失败：{str(e)}")

    else:
        raise HTTPException(status_code=400, detail="不支持的文件格式，请上传 .txt、.md、.docx、.xlsx、.jpg 或 .png")

    # ── Text extraction path: use PromptEngine with Jinja2 template ──
    if extracted_text:
        truncated_text = extracted_text[:4000]
        try:
            response = await prompt_engine.generate(
                template_name="multimodal.j2",
                response_model=MultimodalExtractResponse,
                max_tokens=4000,
                domain=domain,
                extracted_text=truncated_text,
            )
            return response
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"文本抽取失败：{str(e)}")

    raise HTTPException(status_code=400, detail="文档内容为空。")
